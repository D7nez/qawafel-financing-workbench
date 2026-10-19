"""Deterministic evidence calculations. Money is integer halalas throughout."""
from __future__ import annotations
import calendar, csv, hashlib, io, math, re, zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from statistics import median

VERSION = '1.2.0'
START, END = date(2025, 10, 1), date(2026, 9, 30)
MONTHS = [(2025 + (9+i)//12, (9+i)%12+1) for i in range(12)]
FIELDS = ['invoice_id','customer_id','issue_date','due_date','net_amount','vat_amount','currency']

class DataError(ValueError): pass

def money(value):
    """Accept cent values; tolerate only float representation noise <= SAR 1e-7."""
    try:
        if isinstance(value, bool): raise ValueError()
        d = Decimal(str(value).strip())
        if not d.is_finite() or abs(d) >= Decimal('1000000000'): raise ValueError()
        rounded = d.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        if d != rounded:
            if not isinstance(value, float) or abs(d-rounded)>Decimal('0.0000001'):
                raise ValueError()
        return int(rounded*100)
    except (InvalidOperation, ValueError, TypeError):
        raise DataError('Amounts must be finite cent values below SAR 1 billion; only tiny Excel float noise is normalized.')

def iso(value):
    try:
        if isinstance(value, datetime): return value.date()
        if isinstance(value, date): return value
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(value)): raise ValueError()
        return date.fromisoformat(str(value))
    except (ValueError, TypeError):
        raise DataError('Dates must use YYYY-MM-DD.')

def read_invoices(content, filename):
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(x.file_size for x in z.infolist()) > 30_000_000:
                    raise DataError('Expanded workbook exceeds 30 MB.')
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                ws = wb.active
                # Ignore unreliable producer dimension metadata and enforce real streamed limits.
                ws.reset_dimensions()
                header, rows = None, []
                for source_row, values in enumerate(ws.iter_rows(values_only=True),1):
                    if source_row>10001 or len(values)>30:
                        raise DataError('Limit: 10,000 invoice rows and 30 columns.')
                    values=list(values)
                    while values and values[-1] is None: values.pop()
                    if source_row==1:
                        header=[str(x).strip() if x is not None else '' for x in values]
                    elif any(v is not None for v in values):
                        rows.append({**dict(zip(header,values)), '_source_row':source_row})
                if not header: raise DataError('Workbook is empty.')
            finally: wb.close()
        except DataError: raise
        except Exception as e: raise DataError('Cannot read workbook. Upload a valid unencrypted XLSX.') from e
    elif filename.lower().endswith('.csv'):
        try:
            reader = csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
            header = reader.fieldnames or []
            rows = list(reader)
        except Exception as e: raise DataError('Cannot read UTF-8 CSV.') from e
    else: raise DataError('Invoice file must be .xlsx or .csv.')
    if len(rows)>10000: raise DataError('Limit: 10,000 invoice rows.')
    if len(set(header)) != len(header): raise DataError('Duplicate column names are not allowed.')
    if not set(FIELDS).issubset(header): raise DataError('Required columns: '+', '.join(FIELDS))
    return rows

def read_bank(content):
    """Intentionally narrow adapter: generated text PDF, not arbitrary bank layouts."""
    import pdfplumber
    if not content.startswith(b'%PDF-'): raise DataError('Bank statement must be a valid text PDF.')
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if len(pdf.pages)>40: raise DataError('PDF limit: 40 pages.')
            text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
    except DataError: raise
    except Exception as e: raise DataError('Cannot read PDF. Encrypted and scanned PDFs require manual review.') from e
    if 'QAWAFEL_SYNTHETIC_V1' not in text:
        raise DataError('Unsupported PDF layout. Use the supplied synthetic statement; other bank layouts need an adapter/OCR review.')
    statements, txs, seen = [], [], set()
    try:
        for line in text.splitlines():
            parts = [p.strip() for p in line.split('|')]
            if parts[0]=='PERIOD':
                if len(parts)!=5: raise DataError('Malformed statement period.')
                statements.append({'start':iso(parts[1]),'end':iso(parts[2]),'opening':money(parts[3]),'closing':money(parts[4])})
            if parts[0]=='TX':
                if len(parts)!=7: raise DataError('Malformed transaction row.')
                _, tid, dt, kind, ref, amount, curr = parts
                if tid in seen: raise DataError('Duplicate bank transaction ID: '+tid)
                seen.add(tid)
                if curr != 'SAR': raise DataError('Bank statement must use SAR.')
                if kind not in ('RECEIPT','EXPENSE','OWNER_TRANSFER'): raise DataError('Unsupported transaction kind.')
                txs.append({'id':tid,'date':iso(dt),'kind':kind,'reference':ref,'amount':money(amount)})
        expected = [(date(y,m,1),date(y,m,calendar.monthrange(y,m)[1])) for y,m in MONTHS]
        statements.sort(key=lambda x:x['start'])
        if [(s['start'],s['end']) for s in statements] != expected: raise DataError('Bank coverage must be all 12 complete months: Oct 2025-Sep 2026.')
        for i,s in enumerate(statements):
            if i and s['opening'] != statements[i-1]['closing']: raise DataError('Monthly bank balances do not carry forward.')
            change = sum(t['amount'] for t in txs if s['start']<=t['date']<=s['end'])
            if s['opening']+change != s['closing']: raise DataError('Bank balance reconciliation failed for '+str(s['start']))
        if any(not START<=t['date']<=END for t in txs): raise DataError('Bank transaction outside reporting window.')
        if any((t['kind']=='EXPENSE' and t['amount']>=0) or (t['kind']!='EXPENSE' and t['amount']<=0) for t in txs):
            raise DataError('Transaction amount sign does not agree with kind.')
        if not txs: raise DataError('No bank transactions extracted.')
    except DataError: raise
    except Exception as e: raise DataError('Malformed synthetic PDF; review extraction.') from e
    return txs, statements

def analyze(raw, txs, statements, sources=None):
    issues, invoices, seen = [], [], {}
    def issue(level,code,source,message):
        issues.append(dict(level=level,code=code,source=source,message=message))
    for n,row in enumerate(raw,2):
        n=row.get('_source_row',n)
        loc=f'invoices row {n}'
        try:
            if any(row.get(k) is None or str(row.get(k)).strip()=='' for k in FIELDS): raise DataError('Missing required value.')
            if any(str(v).startswith(('=','+','@')) for v in row.values() if v is not None): raise DataError('Formula-like cells are not accepted.')
            iid,cid=str(row['invoice_id']).strip(),str(row['customer_id']).strip()
            if not re.fullmatch(r'[A-Za-z0-9_-]{1,60}',iid) or not re.fullmatch(r'[A-Za-z0-9_-]{1,60}',cid): raise DataError('IDs must be 1-60 letters, numbers, hyphens or underscores.')
            issued,due=iso(row['issue_date']),iso(row['due_date'])
            net,vat=money(row['net_amount']),money(row['vat_amount'])
            if row['currency']!='SAR': raise DataError('Unsupported currency; SAR required.')
            if net<=0 or vat<0: raise DataError('Positive sale and non-negative VAT required; credit notes need a separate adapter.')
            if due<issued or (due-issued).days>365: raise DataError('Due date must be within 0-365 days after issue.')
            if not START<=issued<=END: raise DataError('Issue date outside Oct 2025-Sep 2026.')
            item=dict(id=iid,customer=cid,issued=issued,due=due,net=net,vat=vat,gross=net+vat,row=n)
            fingerprint=(cid,issued,due,net,vat)
            if iid in seen:
                issue('blocker','DUPLICATE_ID',loc,'Duplicate invoice ID '+iid+(' with conflicting values.' if seen[iid]!=fingerprint else '; counted only once.'))
                continue
            seen[iid]=fingerprint
            # Synthetic assumption only, not a tax compliance decision.
            if abs(vat-int((Decimal(net)*Decimal('.15')).quantize(Decimal('1'),rounding=ROUND_HALF_UP)))>1:
                issue('blocker','VAT_MISMATCH',loc,'VAT differs from the synthetic 15% assumption.')
                continue
            invoices.append(item)
        except DataError as e: issue('blocker','INVALID_ROW',loc,str(e))
    coverage={i['issued'].strftime('%Y-%m') for i in invoices}
    months=[f'{y:04}-{m:02}' for y,m in MONTHS]
    missing=[m for m in months if m not in coverage]
    if missing: issue('blocker','MISSING_MONTH','invoices','No valid invoices for '+', '.join(missing)+'. Confirm zero sales or supply data.')
    if not invoices: raise DataError('No valid invoices remain after validation.')
    index={i['id']:i for i in invoices}
    paid=defaultdict(int); matches=[]; unmatched=[]; final_receipt={}
    for t in txs:
        if t['kind']!='RECEIPT': continue
        inv=index.get(t['reference'])
        if inv is None:
            unmatched.append(dict(id=t['id'],reference=t['reference'],amount=t['amount']/100,reason='No valid invoice reference'))
        elif t['date']<inv['issued']:
            unmatched.append(dict(id=t['id'],reference=t['reference'],amount=t['amount']/100,reason='Receipt predates invoice'))
        elif paid[inv['id']]+t['amount']>inv['gross']:
            unmatched.append(dict(id=t['id'],reference=t['reference'],amount=t['amount']/100,reason='Would exceed invoice balance'))
        else:
            paid[inv['id']]+=t['amount']; matches.append(t)
            final_receipt[inv['id']]=max(final_receipt.get(inv['id'],t['date']),t['date'])
    if unmatched: issue('blocker','UNMATCHED_RECEIPTS','bank',f'{len(unmatched)} receipts require allocation review; excluded from matched collections.')
    monthly=[]; customers=defaultdict(int); outstanding=overdue=0; aging={'Not due':0,'1-30 days':0,'31-60 days':0,'61-90 days':0,'90+ days':0}
    for m in months:
        subset=[i for i in invoices if i['issued'].strftime('%Y-%m')==m]
        monthly.append(dict(month=m,revenue=sum(i['net'] for i in subset)/100,
            collections=sum(t['amount'] for t in matches if t['date'].strftime('%Y-%m')==m)/100,
            bank_net_cash=sum(t['amount'] for t in txs if t['date'].strftime('%Y-%m')==m)/100))
    detail=[]
    for i in invoices:
        customers[i['customer']]+=i['net']
        balance=i['gross']-paid[i['id']]; outstanding+=balance
        days=max(0,(END-i['due']).days)
        bucket='Not due' if days==0 else '1-30 days' if days<=30 else '31-60 days' if days<=60 else '61-90 days' if days<=90 else '90+ days'
        aging[bucket]+=balance
        if days>0: overdue+=balance
        detail.append(dict(invoice_id=i['id'],customer=i['customer'],issued=str(i['issued']),due=str(i['due']),net=i['net']/100,gross=i['gross']/100,matched_paid=paid[i['id']]/100,outstanding=balance/100,days_overdue=days if balance else 0,row=i['row']))
    revenue=sum(i['net'] for i in invoices); gross=sum(i['gross'] for i in invoices)
    total_matched=sum(t['amount'] for t in matches)
    receipts=sum(t['amount'] for t in txs if t['kind']=='RECEIPT')
    owner=sum(t['amount'] for t in txs if t['kind']=='OWNER_TRANSFER')
    top=sorted(customers.items(),key=lambda x:(-x[1],x[0]))
    top_pct=100*top[0][1]/revenue
    first=sum(x['revenue'] for x in monthly[:3]); last=sum(x['revenue'] for x in monthly[-3:])
    trend=100*(last/first-1) if first else None
    if top_pct>35: issue('review','CONCENTRATION',top[0][0],f'Top customer represents {top_pct:.1f}% of net invoiced revenue (illustrative review threshold: 35%).')
    overdue_pct=100*overdue/gross
    if overdue_pct>10: issue('review','OVERDUE','invoices',f'{overdue_pct:.1f}% of gross invoiced sales remains overdue (illustrative threshold: 10%).')
    anomalies=model_anomalies(invoices)
    for a in anomalies: a['outstanding']=next(i['gross']-paid[i['id']] for i in invoices if i['id']==a['invoice_id'])/100
    anomalies.sort(key=lambda a:(-a['outstanding'],-a['amount'],a['score'],a['invoice_id']))
    customer_risk=customer_evidence(invoices,paid,final_receipt,revenue)
    overdue90=aging['90+ days']
    overdue90_pct=100*overdue90/overdue if overdue else 0
    no_payment=sum(paid[i['id']]==0 for i in invoices)
    not_overdue=aging['Not due']
    slowest=customer_risk[0]
    for a in anomalies: issue('review','MODEL_ANOMALY',a['invoice_id'],a['explanation'])
    blockers=sum(x['level']=='blocker' for x in issues)
    status='Needs data correction' if blockers else 'Ready for analyst review'
    summary=[f'SAR {overdue90/100:,.2f} is over 90 days overdue, representing {overdue90_pct:.1f}% of all overdue balance. {no_payment} invoices have no matched payment.',
             f"Customer {slowest['id']} has the largest overdue balance: SAR {slowest['overdue']:,.2f}. Review its invoice-level aging before assessing the package.",
             f'Net invoiced revenue is SAR {revenue/100:,.2f} across {len(invoices)} valid invoices and {len(customers)} customers.',
             f'Bank-evidenced collections are SAR {total_matched/100:,.2f}; {100*total_matched/gross:.1f}% of gross invoice value was collected within the reporting window.',
             f'SAR {overdue/100:,.2f} remains overdue as of {END}; this is {overdue_pct:.1f}% of gross invoiced sales.',
             f'Top customer {top[0][0]} contributes {top_pct:.1f}% of net sales. Last-quarter vs first-quarter revenue change: {trend:.1f}%.' if trend is not None else 'Quarter comparison unavailable.',
             f'SAR {owner/100:,.2f} of owner transfers is excluded from collections. {len(unmatched)} receipt(s) need allocation review.',
             f'{blockers} data blocker(s). Package status: {status}. No lending decision or default probability is produced.']
    return dict(version=VERSION,as_of=str(END),period={'start':str(START),'end':str(END)},status=status,
        sources=sources or [],counts={'input_rows':len(raw),'valid_invoices':len(invoices),'bank_transactions':len(txs),'matched_transactions':len(matches),'blockers':blockers},
        metrics={'net_revenue':revenue/100,'gross_invoiced':gross/100,'matched_collections':total_matched/100,'collection_ratio_pct':round(100*total_matched/gross,2),'outstanding':outstanding/100,'overdue':overdue/100,'overdue_pct':round(overdue_pct,2),'top_customer_pct':round(top_pct,2),'quarter_change_pct':round(trend,2) if trend is not None else None,'owner_transfers_excluded':owner/100,'receipt_match_pct':round(100*total_matched/receipts,2) if receipts else None,'closing_cash':statements[-1]['closing']/100,'overdue_90_plus':overdue90/100,'overdue_90_share_pct':round(overdue90_pct,2),'not_overdue_receivables':not_overdue/100,'no_payment_invoice_count':no_payment},
        customer_risk=customer_risk,monthly=monthly,customers=[{'id':k,'revenue':v/100,'share_pct':round(100*v/revenue,2)} for k,v in top],aging={k:v/100 for k,v in aging.items()},issues=issues,unmatched=unmatched,invoices=detail,anomalies=anomalies,summary=summary,
        model={'name':'IsolationForest','seed':42,'features':['log gross amount','payment terms in days','log customer-relative amount'],'training':'Fitted to this upload only; 200 trees, contamination=0.03. Exploratory flags, not fraud labels or credit predictions.','limitations':'Synthetic, unlabeled data; no measured predictive accuracy. A 3% setting forces a tail and is not a measured anomaly rate.'},
        assumptions=['Sales invoices only; one SAR entity; no opening receivables or credit notes.','Fixed closed window Oct 2025-Sep 2026; cutoff Sep 30, 2026.','Net revenue excludes VAT; collections and receivables include VAT.','Exact invoice references are the largest reconciliation simplification. Real exports need amount/date/customer suggestions plus human confirmation; split payments supported.','No cost-of-goods, liabilities or existing debt schedule: affordability and profit cannot be inferred.'])

def customer_evidence(invoices,paid,final_receipt,total_revenue):
    groups=defaultdict(list)
    for i in invoices: groups[i['customer']].append(i)
    result=[]
    for cid,items in groups.items():
        balance=lambda i:i['gross']-paid[i['id']]
        fully_paid=[i for i in items if balance(i)==0]
        paid_gross=sum(i['gross'] for i in fully_paid)
        days=(sum(i['gross']*(final_receipt[i['id']]-i['issued']).days for i in fully_paid)/paid_gross) if paid_gross else None
        result.append({'id':cid,'revenue':sum(i['net'] for i in items)/100,
            'share_pct':round(100*sum(i['net'] for i in items)/total_revenue,2),
            'outstanding':sum(balance(i) for i in items)/100,
            'not_overdue':sum(balance(i) for i in items if i['due']>=END)/100,
            'overdue':sum(balance(i) for i in items if i['due']<END)/100,
            'overdue_90_plus':sum(balance(i) for i in items if (END-i['due']).days>90)/100,
            'invoice_count':len(items),'no_payment_count':sum(paid[i['id']]==0 for i in items),
            'fully_paid_count':len(fully_paid),'weighted_days_to_full_payment':round(days,1) if days is not None else None})
    return sorted(result,key=lambda c:(-c['overdue'],-c['outstanding'],c['id']))

def model_anomalies(invoices):
    from sklearn.ensemble import IsolationForest
    if len(invoices)<30: return []
    group=defaultdict(list)
    for i in invoices: group[i['customer']].append(i['gross'])
    typical={k:median(v) for k,v in group.items()}
    X=[[math.log1p(i['gross']/100),(i['due']-i['issued']).days,math.log1p(i['gross']/typical[i['customer']])] for i in invoices]
    model=IsolationForest(n_estimators=200,contamination=.03,random_state=42,n_jobs=1)
    labels=model.fit_predict(X); scores=model.decision_function(X)
    out=[]
    for i,label,score in zip(invoices,labels,scores):
        if label==-1:
            ratio=i['gross']/typical[i['customer']]
            out.append({'invoice_id':i['id'],'row':i['row'],'score':round(float(score),4),'amount':i['gross']/100,'explanation':f"Unusual amount/terms combination: gross SAR {i['gross']/100:,.2f}, {ratio:.1f}x customer median, {(i['due']-i['issued']).days}-day terms. Check source row {i['row']}; model flag is not proof of error."})
    return sorted(out,key=lambda x:(x['score'],x['invoice_id']))

def digest(content): return hashlib.sha256(content).hexdigest()
