"""Seeded fictional wholesaler. No real names, accounts or customer information."""
import calendar, csv, json, random, sys
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.analysis import MONTHS, END, FIELDS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
ROOT=Path(__file__).resolve().parents[1]; DATA=ROOT/'data'

def generate():
    DATA.mkdir(exist_ok=True)
    rng, payments = random.Random(42), random.Random(142)
    invoices, transactions, planted, monthly_counts = [], [], [], {}
    # Profiles are scenario assumptions, chosen independently of review thresholds.
    profiles = {
        'CUST-01': dict(weight=25, median=6500, terms=[21,30,45], late=[-5,0,2,5,10], unpaid=.02),
        'CUST-02': dict(weight=17, median=4500, terms=[30,45,60], late=[0,3,7,14,21], unpaid=.04),
        'CUST-03': dict(weight=13, median=5500, terms=[30,45,60], late=[25,45,60,90,120], unpaid=.22),
        'CUST-04': dict(weight=11, median=3000, terms=[14,21,30], late=[-3,0,0,3,7], unpaid=.01),
        'CUST-05': dict(weight=9, median=4000, terms=[30,45,60,75], late=[0,5,10,20,30], unpaid=.06),
        'CUST-06': dict(weight=8, median=2800, terms=[14,30,45], late=[0,0,3,7,14], unpaid=.02),
        'CUST-07': dict(weight=6, median=5000, terms=[30,60,90], late=[0,7,14,21,35], unpaid=.05),
        'CUST-08': dict(weight=5, median=2500, terms=[21,30,45], late=[-3,0,0,5,10], unpaid=.02),
        'CUST-09': dict(weight=4, median=6000, terms=[45,60,90,120], late=[45,60,90,120,150], unpaid=.35),
        'CUST-10': dict(weight=2, median=3500, terms=[30,45,60], late=[0,5,10,15,30], unpaid=.05),
    }
    for month_idx,(y,m) in enumerate(MONTHS):
        count = rng.randint(18,30) + (rng.randint(5,10) if m in (2,3) else -3 if m==7 else 0)
        monthly_counts[f'{y}-{m:02}'] = count
        for j in range(count):
            cid=rng.choices(list(profiles),weights=[p['weight'] for p in profiles.values()])[0]
            profile=profiles[cid]
            issued=date(y,m,rng.randint(1,calendar.monthrange(y,m)[1]))
            terms=rng.choice(profile['terms'])
            seasonal=1.12 if m in (2,3) else .90 if m==7 else 1
            # Right-skewed order sizes: many normal orders, occasional bulk orders.
            base=min(60000,max(450,rng.lognormvariate(__import__('math').log(profile['median']),.78)))
            base*=seasonal*(1+month_idx*.009)
            is_planted=month_idx==7 and j==9
            if is_planted:
                cid='CUST-03'; profile=profiles[cid]; base=94000; terms=90
            net=int(Decimal(str(base*100)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
            vat=int((Decimal(net)*Decimal('.15')).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
            iid=f'INV-{y}{m:02}-{j+1:03}'
            invoices.append(dict(zip(FIELDS,[iid,cid,str(issued),str(issued+timedelta(days=terms)),f'{net/100:.2f}',f'{vat/100:.2f}','SAR'])))
            if is_planted:
                planted.append({'invoice_id':iid,'reason':'Deliberately planted large order for a slow-paying customer; sanity check only, not model accuracy evidence.'})
            if not is_planted and payments.random()<profile['unpaid']: continue
            paid_on=issued+timedelta(days=terms+payments.choice(profile['late']))
            gross=net+vat
            if is_planted:
                # A large partially collected order, not an artificially unique payment term.
                parts=[(issued+timedelta(days=60),gross//2)]
            elif payments.random()<.20:
                parts=[(paid_on,gross//2),(paid_on+timedelta(days=payments.randint(7,30)),gross-gross//2)]
            else: parts=[(paid_on,gross)]
            for dt,amt in parts:
                if dt<=END: transactions.append({'date':str(dt),'kind':'RECEIPT','reference':iid,'amount':amt})
    for y,m in MONTHS:
        for day,label,amt in [(5,'RENT',-1250000),(20,'PAYROLL',-2200000),(10,'STOCK_A',-rng.randint(4000000,6500000)),(25,'STOCK_B',-rng.randint(3000000,5000000)),(28,'UTILITIES',-150000)]:
            transactions.append({'date':str(date(y,m,day)),'kind':'EXPENSE','reference':label,'amount':amt})
    for dt in ['2026-01-12','2026-07-12']:
        transactions.append({'date':dt,'kind':'OWNER_TRANSFER','reference':'OWNER_CAPITAL','amount':5000000})
    transactions.sort(key=lambda t:(t['date'],t['reference'],t['amount']))
    for n,t in enumerate(transactions,1): t['id']=f'TX-{n:05}'
    write_invoices(invoices,'invoices_clean')
    dirty=[dict(r) for r in invoices]+[dict(invoices[10])]
    dirty[40]['net_amount']='-200.00'; dirty[80]['currency']='USD'; dirty[120]['due_date']='not-a-date'
    write_invoices(dirty,'invoices_with_issues')
    write_pdf(transactions,'bank_statement.pdf')
    json.dump({'seed':42,'payment_seed':142,'monthly_invoice_counts':monthly_counts,'customer_profiles':profiles,'amount_distribution':'Customer-specific lognormal, sigma 0.78; bounded SAR 450-60,000 before seasonal/growth multipliers; one disclosed planted order.','entity':'Fictional Cedar Wholesale','period':'2025-10-01 to 2026-09-30','invoice_count':len(invoices),'transaction_count':len(transactions),'known_review_cases':planted,'dirty_fixture':'Duplicate final row; negative amount at Excel row 42; USD at row 82; malformed due date at row 122. Corresponding bank receipts may become unmatched.','rounding':'Amounts stored in halalas; synthetic VAT 15%, half-up.','bank':'One fictional account; opening cash SAR 350,000; receipts, expenses and owner capital; no opening receivables.'},open(DATA/'manifest.json','w'),indent=2)
    with open(DATA/'bank_transactions.csv','w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['id','date','kind','reference','amount']);w.writeheader()
        for t in transactions: w.writerow({**t,'amount':f"{t['amount']/100:.2f}"})

def write_invoices(rows,stem):
    with open(DATA/(stem+'.csv'),'w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=FIELDS);w.writeheader();w.writerows(rows)
    wb=Workbook();ws=wb.active;ws.title='Invoices';ws.append(FIELDS)
    for r in rows: ws.append([float(r[k]) if k in ('net_amount','vat_amount') else r[k] for k in FIELDS])
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='123F38')
    for col in 'ABCDEFG':ws.column_dimensions[col].width=23
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(DATA/(stem+'.xlsx'))

def write_pdf(txs,name):
    c=canvas.Canvas(str(DATA/name),pagesize=(595,842),invariant=1)
    balance=35000000
    for y,m in MONTHS:
        start=date(y,m,1);end=date(y,m,calendar.monthrange(y,m)[1])
        subset=[t for t in txs if str(start)<=t['date']<=str(end)]
        closing=balance+sum(t['amount'] for t in subset)
        # Continuation pages repeat layout marker, never duplicate PERIOD metadata.
        chunks=[subset[i:i+42] for i in range(0,len(subset),42)] or [[]]
        for page,chunk in enumerate(chunks):
            c.setFillColor(HexColor('#123F38'));c.setFont('Helvetica-Bold',18);c.drawString(35,800,'SYNTHETIC BANK STATEMENT')
            c.setFont('Helvetica',10);c.drawString(35,778,'Fictional Cedar Wholesale | SAR | Demonstration only')
            c.setFillColor(HexColor('#333333'));c.setFont('Courier',8);c.drawString(35,750,'QAWAFEL_SYNTHETIC_V1')
            if page==0:c.drawString(35,730,f'PERIOD | {start} | {end} | {balance/100:.2f} | {closing/100:.2f}')
            c.drawString(35,708,'TYPE | ID | DATE | KIND | REFERENCE | AMOUNT | CURRENCY')
            pos=688
            for t in chunk:
                c.drawString(35,pos,f"TX | {t['id']} | {t['date']} | {t['kind']} | {t['reference']} | {t['amount']/100:.2f} | SAR");pos-=14
            c.setFont('Helvetica',8);c.drawString(35,35,'All data is fictional. No real bank or account. This is the supported prototype PDF layout.')
            c.showPage()
        balance=closing
    c.save()
if __name__=='__main__':generate();print('Generated variable monthly invoices, XLSX/CSV fixtures and 12-month bank PDF.')
